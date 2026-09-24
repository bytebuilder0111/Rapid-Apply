"""Plain-text extraction from an uploaded resume (PDF or DOCX). The file itself is never
stored; only the AI summary built from this text is."""

import io
import re

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.errors import AppError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
# A typical resume is 3-8k characters; the cap bounds OpenAI cost on unusually long files.
MAX_TEXT_CHARS = 12_000
ALLOWED_EXTENSIONS = (".pdf", ".docx")


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise AppError("resume_unreadable", "This PDF is password-protected.", 400)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise AppError("resume_unreadable", "This PDF file couldn't be read.", 400) from exc


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises assorted zip/xml errors on bad files
        raise AppError("resume_unreadable", "This DOCX file couldn't be read.", 400) from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if not lower.endswith(ALLOWED_EXTENSIONS):
        raise AppError("unsupported_file", "Upload a .pdf or .docx file.", 400)
    if len(data) > MAX_UPLOAD_BYTES:
        raise AppError("file_too_large", "Resume files must be 5 MB or smaller.", 400)

    raw = _pdf_text(data) if lower.endswith(".pdf") else _docx_text(data)
    text = re.sub(r"[ \t]+", " ", raw)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if len(text) < 100:
        raise AppError(
            "resume_unreadable",
            "Couldn't find readable text in this file. If it's a scanned image, upload a "
            "text-based PDF or a DOCX instead.",
            400,
        )
    return text[:MAX_TEXT_CHARS]
